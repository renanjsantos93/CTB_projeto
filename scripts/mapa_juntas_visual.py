"""Mapa de juntas x relatório de inspeção visual de solda (databook BRASNAVAL).

O databook é enviado dividido; o nome de cada parte traz a faixa de páginas do
databook completo (ex.: "VISUAL DE SOLDA - 432 - 462"). A página 1 da parte é
a página inicial da faixa, então:

    página no databook completo = página inicial + (página na parte - 1)

O PDF é só imagem (sem camada de texto); os dados de cada folha foram lidos
visualmente e estão em PARTES abaixo. Para incluir uma nova parte, acrescente
um item em PARTES e rode:

    python scripts/mapa_juntas_visual.py saida/Mapa_Juntas_Visual_Databook.xlsx
"""
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, __file__.rsplit("/", 1)[0])
from padronizar_tag import padronizar_tag  # noqa: E402

# Cada relatório: (página da folha 1/2 na parte, página do isométrico na parte
# ou None, Nº do relatório, data, desenho de referência, juntas do relatório,
# juntas numeradas no isométrico ou None quando não conferido, observação).
PARTES = [
    {
        "arquivo": "VISUAL DE SOLDA - 432 - 462.pdf",
        "inicio": 432,
        "fim": 462,
        "paginas_pdf": 31,
        "relatorios": [
            (1, 2, "001/26", "14/08/2026", '1"-A1E-CF-4310A',
             ["J-1", "J-2", "J-3", "J-4", "J-5", "J-6"], range(1, 7), ""),
            (3, 4, "001/26", "17/08/2026", '1"-B1E-MR-4005A',
             ["J-1", "J-2", "J-3"], range(1, 5),
             'J-4 = ligação com 2"-B1E-MR-4010A: conferir se está no relatório da 4010A'),
            (5, 6, "001/26", "17/08/2026", '1"-B1E-MR-4006A',
             ["J-1", "J-2", "J-3"], range(1, 5),
             'J-4 = ligação com 2"-B1E-MR-4010A: conferir se está no relatório da 4010A'),
            (7, 8, "001/26", "17/08/2026", '1"-B1E-MR-4007A',
             ["J-1", "J-2", "J-3"], range(1, 5),
             'J-4 = ligação com 2"-B1E-MR-4010A: conferir se está no relatório da 4010A'),
            (9, 10, "001/26", "17/08/2026", '1"-B1E-MR-4008A',
             ["J-1", "J-2", "J-3"], range(1, 5),
             'J-4 = ligação com 2"-B1E-MR-4010A: conferir se está no relatório da 4010A'),
            (11, 12, "001/26", "17/08/2026", '1"-B1E-MR-4009A',
             ["J-1", "J-2", "J-3"], range(1, 5),
             'J-4 = ligação com 2"-B1E-MR-4010A: conferir se está no relatório da 4010A'),
            (13, 14, "001/26", "17/08/2026", '1"-B3B-LO-5003A',
             ["J-1", "J-2", "J-3", "J-5", "J-6", "J-7", "J-8", "J-9", "J-10", "J-11"],
             range(1, 12), ""),
            (15, 16, "001/26", "17/08/2026", '1"-D1E-LPG-4305A-C-3',
             ["J-1", "J-2", "J-3", "J-4", "J-5", "J-6"], range(1, 8), ""),
            (17, 18, "001/26", "18/08/2026", '1"-D1E-LPG-4308A-C-3',
             ["J-%d" % i for i in range(1, 10)], range(1, 10), ""),
            (19, 20, "001/26", "18/08/2026", '1"-D1E-NG-4001A FL 1/2',
             ["J-%d" % i for i in range(1, 13)], None, ""),
            (21, 22, "001/26", "18/08/2026", '1"-D1E-NG-4001A FL 2/2',
             ["J-%d" % i for i in range(12, 22)], None, ""),
            (23, 24, "001/26", "18/08/2026", '2"-A1E-CF-4007A',
             ["J-02", "J-04"], range(1, 5), ""),
            (25, 26, "001/26", "18/08/2026", '2"-B1E-MR-4010A FL 1/2',
             ["J-1", "J-2", "J-4", "J-8"], range(1, 10),
             "Folha com os tês dos ramais 1\"-B1E-MR-4006A a 4009A: conferir se a junta está em outro relatório"),
            (27, 28, "001/26", "18/08/2026", '2"-B1E-MR-4010A FL 2/2',
             ["J-10", "J-12", "J-13", "J-17", "J-18", "J-19"], range(10, 20),
             "J-14/J-15/J-16 junto às válvulas: verificar se são juntas flangeadas"),
            (29, 30, "001/26", "18/08/2026", '2"-D1E-LPG-4002A-C-3',
             ["J-2", "J-3", "J-4"], range(1, 6), ""),
            (31, None, "001/26", "22/08/2026", '1"-A3B-N2-9105A',
             ["J-1", "J-2", "J-5", "J-6", "J-7", "J-8", "J-11", "J-12", "J-13",
              "J-14", "J-15"], None,
             "Folha 2/2 (isométrico) na pág. 463, próxima parte do databook; numeração pula J-3, J-4, J-9, J-10: conferir no isométrico"),
        ],
    },
]

COLUNAS = [
    ("TAG PADRONIZADA", 24), ("LINHA", 24), ("DOCUMENTO", 30), ("JUNTA", 9),
    ("Nº DO RELATÓRIO", 15), ("PÁGINA DO DATABOOK", 14),
    ("PÁGINA DO ISOMÉTRICO", 14), ("DATA DO RELATÓRIO", 13),
    ("PÁGINA NO ARQUIVO DA PARTE", 14), ("ARQUIVO DA PARTE", 30),
    ("CONFERÊNCIA", 60),
]


def numero_junta(j):
    return int(re.sub(r"\D", "", j))


def linha_do_documento(doc):
    return re.sub(r"\s*FL\s*\d+\s*/\s*\d+\s*$", "", doc, flags=re.I).strip()


def montar_linhas():
    linhas, controle = [], []
    vistos = {}
    for parte in PARTES:
        ini, fim = parte["inicio"], parte["fim"]
        esperado = fim - ini + 1
        controle.append((parte["arquivo"], ini, fim, esperado, parte["paginas_pdf"],
                         "OK" if esperado == parte["paginas_pdf"] else
                         "DIVERGENTE: nº de páginas não bate com a faixa"))
        for (pg, pg_iso, rel, data, doc, juntas, iso, obs) in parte["relatorios"]:
            linha = linha_do_documento(doc)
            tag, obs_tag = padronizar_tag(linha)
            pag_db = ini + pg - 1
            pag_iso = ini + pg_iso - 1 if pg_iso else None
            no_rel = {numero_junta(j) for j in juntas}
            for j in juntas:
                n = numero_junta(j)
                conf = []
                if obs_tag:
                    conf.append(obs_tag)
                if iso is not None and n not in iso:
                    conf.append("Junta não aparece no isométrico")
                chave = (tag, n)
                if chave in vistos:
                    conf.append("Junta repetida: também na pág. %d" % vistos[chave])
                else:
                    vistos[chave] = pag_db
                linhas.append([tag, linha, doc, "%03d" % n, rel, pag_db, pag_iso,
                               data, pg, parte["arquivo"],
                               "; ".join(conf) or "OK"])
            if iso is not None:
                for n in iso:
                    if n in no_rel:
                        continue
                    conf = "FALTA NO RELATÓRIO: junta numerada no isométrico (pág. %d) " \
                           "sem registro de inspeção visual" % pag_iso
                    if obs:
                        conf += " | " + obs
                    linhas.append([tag, linha, doc, "%03d" % n, "", pag_db, pag_iso,
                                   data, pg, parte["arquivo"], conf])
            elif obs:
                linhas[-1][-1] = (linhas[-1][-1] + "; " if linhas[-1][-1] != "OK" else "") + obs
    return linhas, controle


def gravar(saida):
    linhas, controle = montar_linhas()
    linhas.sort(key=lambda r: (r[5], r[2], r[3]))

    fino = Side(style="thin", color="999999")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    cab_fill = PatternFill("solid", fgColor="5B7A5F")
    cab_font = Font(bold=True, color="FFFFFF")
    alerta = PatternFill("solid", fgColor="F8CBAD")
    aviso = PatternFill("solid", fgColor="FFE699")

    wb = Workbook()
    ws = wb.active
    ws.title = "MAPA DE JUNTAS"
    for c, (nome, larg) in enumerate(COLUNAS, 1):
        cel = ws.cell(1, c, nome)
        cel.fill, cel.font, cel.border = cab_fill, cab_font, borda
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = larg
    ws.row_dimensions[1].height = 32
    for r, dados in enumerate(linhas, 2):
        for c, v in enumerate(dados, 1):
            cel = ws.cell(r, c, v)
            cel.border = borda
            cel.alignment = Alignment(horizontal="left" if c in (3, 10, 11) else "center",
                                      vertical="center", wrap_text=(c == 11))
            if c == 4:
                cel.number_format = "@"
        conf = dados[-1]
        if conf.startswith("FALTA"):
            for c in range(1, len(COLUNAS) + 1):
                ws.cell(r, c).fill = alerta
        elif conf != "OK":
            ws.cell(r, len(COLUNAS)).fill = aviso
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(COLUNAS)), len(linhas) + 1)

    # Relatórios: um por folha 1/2
    wr = wb.create_sheet("RELATÓRIOS")
    cab = ["PÁGINA DO DATABOOK", "PÁGINA DO ISOMÉTRICO", "Nº DO RELATÓRIO", "DATA",
           "TAG PADRONIZADA", "DOCUMENTO", "QTD. JUNTAS NO RELATÓRIO",
           "QTD. JUNTAS NO ISOMÉTRICO", "JUNTAS FALTANDO NO RELATÓRIO", "CONFERÊNCIA"]
    larg = [12, 12, 13, 12, 24, 30, 12, 12, 22, 60]
    for c, (n, w) in enumerate(zip(cab, larg), 1):
        cel = wr.cell(1, c, n)
        cel.fill, cel.font, cel.border = cab_fill, cab_font, borda
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        wr.column_dimensions[get_column_letter(c)].width = w
    wr.row_dimensions[1].height = 32
    contagem_rel = {}
    for parte in PARTES:
        for rel in parte["relatorios"]:
            contagem_rel[rel[2]] = contagem_rel.get(rel[2], 0) + 1
    r = 2
    for parte in PARTES:
        ini = parte["inicio"]
        for (pg, pg_iso, rel, data, doc, juntas, iso, obs) in parte["relatorios"]:
            nums = {numero_junta(j) for j in juntas}
            falt = [n for n in iso if n not in nums] if iso is not None else []
            conf = []
            if contagem_rel[rel] > 1:
                conf.append("Nº DE RELATÓRIO REPETIDO (%s usado em %d relatórios)"
                            % (rel, contagem_rel[rel]))
            if iso is None:
                conf.append("Isométrico não conferido junta a junta")
            if obs:
                conf.append(obs)
            vals = [ini + pg - 1, ini + pg_iso - 1 if pg_iso else None, rel, data,
                    padronizar_tag(linha_do_documento(doc))[0], doc, len(juntas),
                    len(iso) if iso is not None else None,
                    ", ".join("%03d" % n for n in falt), "; ".join(conf)]
            for c, v in enumerate(vals, 1):
                cel = wr.cell(r, c, v)
                cel.border = borda
                cel.alignment = Alignment(horizontal="left" if c in (6, 10) else "center",
                                          vertical="center", wrap_text=c in (9, 10))
            if falt:
                wr.cell(r, 9).fill = alerta
            r += 1
    wr.freeze_panes = "A2"

    # Controle de partes
    wc = wb.create_sheet("CONTROLE DE PARTES")
    cab = ["ARQUIVO", "PÁGINA INICIAL", "PÁGINA FINAL", "PÁGINAS ESPERADAS",
           "PÁGINAS NO PDF", "SITUAÇÃO"]
    for c, (n, w) in enumerate(zip(cab, [34, 14, 14, 14, 14, 40]), 1):
        cel = wc.cell(1, c, n)
        cel.fill, cel.font, cel.border = cab_fill, cab_font, borda
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        wc.column_dimensions[get_column_letter(c)].width = w
    for r, vals in enumerate(controle, 2):
        for c, v in enumerate(vals, 1):
            cel = wc.cell(r, c, v)
            cel.border = borda
            cel.alignment = Alignment(horizontal="center")
    r = len(controle) + 3
    cobertas = sum(p["fim"] - p["inicio"] + 1 for p in PARTES)
    wc.cell(r, 1, "Páginas do databook conferidas: %d de 856" % cobertas).font = Font(bold=True)
    wc.cell(r + 1, 1, "Regra: página no databook = página inicial + (página na parte - 1)")

    wb.save(saida)
    return linhas


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    ls = gravar(sys.argv[1])
    faltas = sum(1 for l in ls if l[-1].startswith("FALTA"))
    print("%d linhas (%d juntas faltando no relatório) -> %s" % (len(ls), faltas, sys.argv[1]))
