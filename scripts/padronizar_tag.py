"""Padronização de TAG de linha nos mapas de juntas.

Formato alvo: DN-CLASS-SERVICE-LINE_NUMBER  (ex.: 1.1/2"-A1E-RA-2302)

Regras:
 1. Remove espaços.
 2. Ignora maiúsculas/minúsculas (saída em maiúsculas).
 3. Símbolos equivalentes a polegada (' ° º ’’ ″ ” “ ...) viram ".
 4. Corrige variações de OCR do DN 1.1/2" (1-1/2", 1-1/2, 1-12, 1.12).
    4-1/2" só é aceito como 1.1/2" se a Line List confirmar; caso contrário
    é mantido como 4.1/2" e sinalizado para validação.
 5. Mantém somente os 4 primeiros blocos.

Uso:
    python scripts/padronizar_tag.py ENTRADA.xlsx SAIDA.xlsx [--line-list LL.xlsx|LL.txt]

A planilha é alterada diretamente no XML da aba (apenas as células da coluna
"TAG PADRONIZADA"), preservando logos, comentários, formatação e fórmulas.
"""
import argparse
import re
import sys
import zipfile
from xml.sax.saxutils import escape

import openpyxl

INCH_CHARS = "'°º’‘″”“ʺ˝″‶`´"
DASHES = "‐‑‒–—―−"

# DN 1.1/2" digitado/lido por OCR de formas diferentes (somente no 1º bloco)
_OCR_1_1_2 = re.compile(r'^1(?:[-.]1/2|[-.]12)"?(?=-|$)')
_FRACAO_HIFEN = re.compile(r'^(\d+)-(\d+/\d+)"?(?=-|$)')
_DN_SEM_POL = re.compile(r'^(\d+(?:\.\d+/\d+|/\d+)?)(?=-|$)')


def _limpar(tag):
    s = str(tag).upper()
    s = re.sub(r"\s+", "", s)
    for ch in DASHES:
        s = s.replace(ch, "-")
    s = re.sub("[" + re.escape(INCH_CHARS) + "]+", '"', s)
    s = re.sub(r'"{2,}', '"', s)
    return s


def padronizar_tag(tag, line_list=None):
    """Retorna (tag_padronizada, observacao). observacao é '' quando ok."""
    if tag is None or str(tag).strip() == "":
        return "", ""
    s = _limpar(tag)
    obs = ""

    s = _OCR_1_1_2.sub('1.1/2"', s)

    m = _FRACAO_HIFEN.match(s)
    if m and m.group(1) == "4" and m.group(2) == "1/2":
        resto = s[m.end():]
        candidato_1 = _quatro_blocos('1.1/2"' + resto)
        candidato_4 = _quatro_blocos('4.1/2"' + resto)
        if line_list and candidato_1 in line_list and candidato_4 not in line_list:
            s = '1.1/2"' + resto
            obs = 'DN 4-1/2" corrigido para 1.1/2" conforme Line List'
        else:
            s = '4.1/2"' + resto
            obs = 'VALIDAR DN 4-1/2" CONTRA LINE LIST (possível 1.1/2")'
    elif m:
        s = _FRACAO_HIFEN.sub(r'\1.\2"', s)

    # DN sem símbolo de polegada (ex.: 2-A1E-...) recebe "
    s = _DN_SEM_POL.sub(r'\1"', s)

    s = _quatro_blocos(s)
    if s.count("-") < 3:
        obs = (obs + "; " if obs else "") + "TAG COM MENOS DE 4 BLOCOS"
    return s, obs


def _quatro_blocos(s):
    blocos = [b for b in s.split("-") if b != ""]
    return "-".join(blocos[:4])


def carregar_line_list(caminho):
    if not caminho:
        return None
    tags = set()
    if caminho.lower().endswith((".xlsx", ".xlsm")):
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str) and v.count("-") >= 3:
                        tags.add(padronizar_tag(v)[0])
    else:
        with open(caminho, encoding="utf-8") as f:
            for linha in f:
                if linha.strip():
                    tags.add(padronizar_tag(linha)[0])
    return tags


# --------------------------------------------------------------------------
# Escrita no xlsx preservando o restante do arquivo
# --------------------------------------------------------------------------

def _norm_header(v):
    return re.sub(r"\s+", " ", str(v or "")).strip().upper()


CABECALHOS_ORIGEM = {
    "LINHA", "TAG", "TAG LINHA", "TAG DA LINHA", "LINE", "LINE NUMBER",
    "Nº DA LINHA", "N° DA LINHA", "NO DA LINHA", "N DA LINHA",
}


def localizar_colunas(ws, max_linhas=30):
    """Encontra a coluna TAG PADRONIZADA, a coluna de origem e a 1ª linha de dados."""
    destino = origem = None
    linha_hdr = 0
    for row in ws.iter_rows(min_row=1, max_row=max_linhas):
        for c in row:
            h = _norm_header(c.value)
            if h == "TAG PADRONIZADA":
                destino, linha_hdr = c.column, max(linha_hdr, c.row)
            elif origem is None and re.sub(r"\d+$", "", h) in CABECALHOS_ORIGEM:
                origem, linha_hdr = c.column, max(linha_hdr, c.row)
    return destino, origem, linha_hdr + 1


def _col_letra(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _caminho_aba(zf, nome_aba):
    wb_xml = zf.read("xl/workbook.xml").decode("utf-8")
    rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    m = re.search(r'<sheet [^>]*name="%s"[^>]*r:id="([^"]+)"' % re.escape(escape(nome_aba, {'"': "&quot;"})), wb_xml)
    rid = m.group(1)
    alvo = re.search(r'<Relationship [^>]*Id="%s"[^>]*Target="([^"]+)"' % rid, rels) or \
        re.search(r'<Relationship [^>]*Target="([^"]+)"[^>]*Id="%s"' % rid, rels)
    t = alvo.group(1).lstrip("/")
    return t if t.startswith("xl/") else "xl/" + t


def _gravar_celula(xml, ref, texto):
    val = '<is><t>%s</t></is>' % escape(texto)
    # célula já existente (vazia ou com valor)
    pad = re.compile(r'<c r="%s"((?:\s+[a-zA-Z:]+="[^"]*")*)\s*(?:/>|>(?:(?!<c[ >]).)*?</c>)' % ref, re.S)
    if pad.search(xml):
        def _sub(m):
            attrs = re.sub(r'\s+t="[^"]*"', "", m.group(1))
            return '<c r="%s"%s t="inlineStr">%s</c>' % (ref, attrs, val)
        return pad.sub(_sub, xml, count=1)
    # célula inexistente: insere na linha, respeitando a ordem das colunas
    lin = re.match(r"[A-Z]+(\d+)", ref).group(1)
    col = re.match(r"([A-Z]+)", ref).group(1)
    nova = '<c r="%s" t="inlineStr">%s</c>' % (ref, val)
    mrow = re.search(r'(<row r="%s"[^>]*?)(/>|>(.*?)</row>)' % lin, xml, re.S)
    if not mrow:
        raise RuntimeError("linha %s não encontrada no XML" % lin)
    if mrow.group(2) == "/>":
        return xml[:mrow.start()] + mrow.group(1) + ">" + nova + "</row>" + xml[mrow.end():]
    corpo = mrow.group(3)
    chave = (len(col), col)
    pos = len(corpo)
    for mc in re.finditer(r'<c r="([A-Z]+)\d+"', corpo):
        if (len(mc.group(1)), mc.group(1)) > chave:
            pos = mc.start()
            break
    corpo = corpo[:pos] + nova + corpo[pos:]
    ini = mrow.start(3)
    return xml[:ini] + corpo + xml[mrow.end(3):]


def processar(entrada, saida, aba=None, line_list=None):
    wb = openpyxl.load_workbook(entrada, read_only=False, data_only=False)
    ws = wb[aba] if aba else wb.worksheets[0]
    destino, origem, primeira = localizar_colunas(ws)
    if not destino or not origem:
        raise SystemExit("Colunas 'TAG PADRONIZADA' e/ou 'Linha' não encontradas na aba %r" % ws.title)

    resultados = {}
    pendencias = []
    for r in range(primeira, ws.max_row + 1):
        bruto = ws.cell(r, origem).value
        if bruto is None or str(bruto).strip() == "":
            continue
        tag, obs = padronizar_tag(bruto, line_list)
        resultados["%s%d" % (_col_letra(destino), r)] = tag
        if obs:
            pendencias.append((r, bruto, tag, obs))

    with zipfile.ZipFile(entrada) as zin:
        caminho = _caminho_aba(zin, ws.title)
        xml = zin.read(caminho).decode("utf-8")
        for ref, tag in resultados.items():
            xml = _gravar_celula(xml, ref, tag)
        with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                dados = xml.encode("utf-8") if item.filename == caminho else zin.read(item.filename)
                zout.writestr(item, dados)

    return ws.title, resultados, pendencias


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("entrada")
    ap.add_argument("saida")
    ap.add_argument("--aba", help="nome da aba (padrão: primeira)")
    ap.add_argument("--line-list", help="Line List (.xlsx ou .txt, uma TAG por linha) para validar 4-1/2\"")
    a = ap.parse_args(argv)

    titulo, res, pend = processar(a.entrada, a.saida, a.aba, carregar_line_list(a.line_list))
    print("Aba %s: %d TAGs padronizadas, %d únicas" % (titulo, len(res), len(set(res.values()))))
    for r, bruto, tag, obs in pend:
        print("  linha %d: %r -> %r  [%s]" % (r, bruto, tag, obs))
    return 0


if __name__ == "__main__":
    sys.exit(main())
