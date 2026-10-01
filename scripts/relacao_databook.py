"""Relação de relatórios do databook BRASNAVAL (relatórios dimensionais Master).

Extrai do PDF (camada de texto) o nº, tipo e data de cada relatório, o desenho
de referência e as cotas medidas; as juntas vêm da leitura dos balões do
croqui de cada relatório (tabela JUNTAS_CROQUI abaixo), pois só existem como
imagem no PDF.

Uso:
    python scripts/relacao_databook.py DATABOOK.pdf SAIDA.xlsx
"""
import re
import subprocess
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, __file__.rsplit("/", 1)[0])
from padronizar_tag import padronizar_tag  # noqa: E402

# Juntas lidas nos balões do croqui, por nº de página do databook.
# None = balões ilegíveis (imagem de baixa resolução no PDF).
JUNTAS_CROQUI = {
    1: list(range(4, 16)),
    2: list(range(5, 16)) + [17],
    3: [1, 2, 3],
    4: None,
    5: [14, 15, 16, 17],
    6: list(range(16, 25)),
    7: list(range(15, 21)),
    8: [6, 7, 8, 9],
    9: list(range(1, 8)),
    10: list(range(38, 51)),
    11: list(range(5, 14)),
    12: None,
    13: None,
    14: [1, 2, 3, 4],
    15: [5, 6, 7, 8],
    16: list(range(15, 21)),
    17: [14, 15, 16, 17],
    18: [2, 3, 4, 5],
    19: [6, 7, 8, 9],
    20: list(range(1, 8)),
    21: list(range(5, 16)) + [17],
    22: None,
    23: list(range(1, 12)),
    24: None,
    25: None,
    26: None,
    27: None,
    28: list(range(1, 9)),
    29: list(range(16, 25)),
    30: None,
    31: list(range(5, 13)),
    32: list(range(21, 29)),
    33: None,
    34: list(range(4, 16)),
    35: list(range(13, 21)),
    36: list(range(5, 16)) + [17],
}

# Spools e observações escritas no croqui.
OBS_CROQUI = {
    1: ("01, 02", "Divisão entre spools na junta 8; juntas 4 e 15 solda de campo; spool 02 com sobremetal"),
    2: ("01 a 04", "Spools divididos nas juntas 13, 11 e 09"),
    3: ("01", ""),
    4: ("", ""),
    5: ("01", "Extremidades com sobremetal"),
    6: ("01, 02", "Spool dividido na junta 19; extremidade 01 com sobremetal"),
    7: ("", "Extremidade com sobremetal"),
    8: ("", ""),
    9: ("01, 02", "Spool dividido na junta 4"),
    10: ("01", ""),
    11: ("01, 02", "Spool dividido na junta 8"),
    12: ("", ""),
    13: ("", ""),
    14: ("", "Extremidades FT (solda de campo)"),
    15: ("", "Extremidade FT (solda de campo)"),
    16: ("", "Extremidade com sobremetal"),
    17: ("01", "Extremidades com sobremetal"),
    18: ("01", ""),
    19: ("", ""),
    20: ("01, 02", "Spool dividido na junta 4"),
    21: ("01 a 04", "Spools divididos nas juntas 13, 11 e 09"),
    22: ("", ""),
    23: ("01", ""),
    24: ("01 a 03", "Spool dividido nas juntas 06 e 20"),
    25: ("01", ""),
    26: ("01 a 03", "Spool dividido nas juntas 04 e 12"),
    27: ("01, 02", "Spool dividido na junta 04"),
    28: ("01 a 03", "Spools divididos nas juntas 04 e 05; extremidades FT"),
    29: ("01, 02", "Spool dividido na junta 19; extremidade 01 com sobremetal"),
    30: ("01 a 04", "Spool dividido nas juntas 08, 16 e 28"),
    31: ("01 a 03", "Spools divididos nas juntas 08 e 09; extremidades FT"),
    32: ("01 a 03", "Spools divididos nas juntas 24 e 25; extremidades FT"),
    33: ("02, 03", "Spool dividido nas juntas 08 e 12"),
    34: ("01, 02", "Spool dividido na junta 8; juntas 4 e 15 FW"),
    35: ("01 a 03", "Spools divididos nas juntas 16 e 17; extremidades FT"),
    36: ("01 a 04", "Spools divididos nas juntas 13, 11 e 09"),
}

# Relatórios cujo croqui é idêntico ao de outro relatório (mesma página-base).
CROQUI_IGUAL = {16: 7, 17: 5, 19: 8, 20: 9, 21: 2, 29: 6, 34: 1, 36: 2}


def extrair_paginas(pdf):
    txt = subprocess.run(["pdftotext", "-layout", pdf, "-"], check=True,
                         capture_output=True, text=True).stdout
    return [p for p in txt.split("\f") if p.strip()]


def tag_do_desenho(desenho):
    s = re.sub(r"\s*[-/]?\s*FOLHA.*$", "", desenho, flags=re.I).strip()
    if s.upper().startswith("I-IS-"):
        return ""  # isométrico, não traz a TAG da linha
    s = re.sub(r"^MO\d+-", "", s, flags=re.I)
    return padronizar_tag(s)[0]


def ler_relatorio(pag, p):
    num = re.search(r"Seq:\s*(\d+)", p).group(1)
    data = re.search(r"DATA:\s*(\d{2}/\d{2}/\d{4})", p).group(1)
    tipo = re.search(r"(RELAT[ÓO]RIO [A-ZÇÃÕÉÍÓÚ ]+?)\s{2,}", p).group(1).strip()
    m = re.search(r"DRAWING\s+REV\.\s*\n(.*)", p).group(1)
    campos = re.split(r"\s{2,}", m.strip())
    desenho, rev_des = campos[-2], campos[-1]
    cotas = re.findall(r"^\s*(\d{2})\s+(\d+)\s+Conforme norma\s+(\d+)", p, re.M)
    return {
        "pagina": pag, "numero": num, "tipo": tipo.title(),
        "data": datetime.strptime(data, "%d/%m/%Y").date(),
        "desenho": desenho, "rev_desenho": rev_des, "procedimento": campos[1],
        "tag": tag_do_desenho(desenho), "cotas": cotas,
        # caixa de seleção é imagem; todos os croquis conferidos estão "APROVADO"
        "laudo": "Aprovado",
    }


def obs_relatorio(r, rels):
    pag = r["pagina"]
    obs = []
    outros = [o["numero"] for o in rels
              if o["desenho"] == r["desenho"] and o["pagina"] != pag]
    if outros:
        obs.append("Mesmo desenho dos relatórios " + ", ".join(outros))
    if not r["tag"]:
        obs.append("TAG não consta no relatório (referência é o isométrico)")
    if JUNTAS_CROQUI[pag] is None:
        obs.append("Nº das juntas ilegível no croqui – conferir no isométrico")
    if pag in CROQUI_IGUAL:
        obs.append(f"Croqui idêntico ao do relatório {CROQUI_IGUAL[pag]:03d}")
    if OBS_CROQUI[pag][1]:
        obs.append("Croqui: " + OBS_CROQUI[pag][1])
    return "; ".join(obs)


HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF")
ALERTA = PatternFill("solid", fgColor="FFF2CC")
FINO = Side(style="thin", color="999999")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)


def escrever_aba(ws, cabecalho, linhas, larguras, alerta_col=None):
    ws.append(cabecalho)
    for c in ws[1]:
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDA
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for linha in linhas:
        ws.append(linha)
        for c in ws[ws.max_row]:
            c.border = BORDA
            c.alignment = Alignment(vertical="center", wrap_text=True)
            if hasattr(c.value, "year"):
                c.number_format = "DD/MM/YYYY"
                c.alignment = Alignment(horizontal="center", vertical="center")
        if alerta_col is not None and linha[alerta_col]:
            ws.cell(ws.max_row, alerta_col + 1).fill = ALERTA
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main(pdf, saida):
    rels = [ler_relatorio(i, p) for i, p in enumerate(extrair_paginas(pdf), 1)]
    wb = Workbook()

    # 1) Uma linha por junta
    linhas, item = [], 0
    for r in rels:
        juntas = JUNTAS_CROQUI[r["pagina"]] or [""]
        for j in juntas:
            item += 1
            linhas.append([item, r["numero"], r["tipo"], r["data"], r["desenho"],
                           r["tag"], f"{j:02d}" if j != "" else "",
                           OBS_CROQUI[r["pagina"]][0], r["laudo"], r["pagina"],
                           obs_relatorio(r, rels)])
    ws = wb.active
    ws.title = "Relação por junta"
    escrever_aba(ws, ["Item", "Nº do relatório", "Tipo de relatório", "Data do relatório",
                      "Documento (desenho de referência)", "TAG da linha", "Nº da junta",
                      "Spool(s)", "Laudo", "Página no databook", "Observação"],
                 linhas, [6, 11, 20, 12, 40, 22, 9, 10, 10, 10, 70], alerta_col=10)

    # 2) Uma linha por relatório
    linhas = []
    for r in rels:
        js = JUNTAS_CROQUI[r["pagina"]]
        linhas.append([r["numero"], r["tipo"], r["data"], r["desenho"], r["rev_desenho"],
                       r["tag"], ", ".join(f"{j:02d}" for j in js) if js else "ILEGÍVEL",
                       len(js) if js else "", OBS_CROQUI[r["pagina"]][0],
                       r["procedimento"], len(r["cotas"]), r["laudo"], r["pagina"],
                       obs_relatorio(r, rels)])
    ws = wb.create_sheet("Resumo por relatório")
    escrever_aba(ws, ["Nº do relatório", "Tipo de relatório", "Data do relatório",
                      "Documento (desenho de referência)", "Rev. desenho", "TAG da linha",
                      "Juntas no croqui", "Qtd. juntas", "Spool(s)", "Procedimento",
                      "Qtd. cotas", "Laudo", "Página no databook", "Observação"],
                 linhas, [11, 20, 12, 40, 8, 22, 40, 8, 10, 14, 8, 10, 10, 70],
                 alerta_col=13)

    # 3) Cotas com valor encontrado diferente do especificado
    linhas = []
    for r in rels:
        obs_c = OBS_CROQUI[r["pagina"]][1]
        sobremetal = "sobremetal" in obs_c.lower() or "FT" in obs_c
        for it, esp, enc in r["cotas"]:
            d = int(enc) - int(esp)
            if d == 0:
                continue
            if abs(d) <= 3:
                av = "Diferença pequena (≤ 3 mm)"
            elif d == 100 and sobremetal:
                av = "+100 mm – croqui indica sobremetal/FT"
            else:
                av = "VERIFICAR – laudo marcado como 'Atende'"
            linhas.append([r["numero"], r["desenho"], r["tag"], it, int(esp), int(enc), d,
                           av, r["pagina"]])
    ws = wb.create_sheet("Divergências de cota")
    escrever_aba(ws, ["Nº do relatório", "Documento (desenho de referência)", "TAG da linha",
                      "Item da cota", "Especificado (mm)", "Encontrado (mm)",
                      "Diferença (mm)", "Avaliação", "Página no databook"],
                 linhas, [11, 40, 22, 9, 13, 13, 12, 42, 10])
    for row in ws.iter_rows(min_row=2):
        if row[7].value.startswith("VERIFICAR"):
            row[7].fill = ALERTA

    wb.save(saida)
    print(f"{len(rels)} relatórios -> {saida}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
