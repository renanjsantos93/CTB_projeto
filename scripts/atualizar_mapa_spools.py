"""Atualiza as etapas de fabricação do Mapa de Spools a partir do Mapa de Juntas.

Regra (procedimento "Atualização Automática do Mapa de Spools"):
 1. Só entram na análise as juntas Pipe Shop (P/C = P / PIPE / PIPE SHOP).
    Juntas de Campo (P/C = C / CAMPO) e sem P/C são desconsideradas.
 2. Uma etapa do spool só é preenchida quando TODAS as juntas Pipe Shop do
    spool têm data na etapa correspondente do Mapa de Juntas.
 3. A data gravada é a maior (mais recente) entre essas juntas.
 4. Se pelo menos uma junta Pipe Shop está sem data, o campo do spool fica
    em branco (valores antigos são apagados, salvo com --nao-limpar).

Correspondência de etapas (Mapa de Juntas -> Mapa de Spools):
    Visual de Ajuste (VA) -> Data de Corte / VA
    Ajuste                -> Data de Ajuste
    Soldagem              -> Data de Soldagem
    EVS (Ensaio Visual)   -> Data de EVS
    END (LP/PM, RX/US)    -> Data de END
    Dimensional           -> Data Dimensional
Etapas sem coluna em um dos mapas são ignoradas e avisadas no relatório.

END: a data da junta é a mais recente entre os ensaios (LP, PM, RX, US...).
Ensaios com status "N" (não aplicável) são ignorados; um ensaio com status
"P" (pendente) e sem data deixa a junta sem data de END.

A associação junta -> spool é feita por Linha (com zona, ex.: -Z02) + nº do
spool. Se a linha do Mapa de Spools não tiver zona, usa-se a TAG padronizada
(4 blocos) quando ela identifica uma única linha do Mapa de Juntas.

Uso:
    python scripts/atualizar_mapa_spools.py MAPA_JUNTAS.xlsx MAPA_SPOOLS.xlsx SAIDA.xlsx \\
        [--aba-juntas "SGJ (2)"] [--aba-spools "Mapa de Spools"] \\
        [--criar-colunas] [--nao-limpar] [--relatorio RELATORIO.csv]

Só as células das colunas de data do Mapa de Spools são alteradas; o restante
do arquivo (logos, fórmulas, formatação) é preservado.
"""
import argparse
import csv
import datetime as dt
import re
import sys
import unicodedata
import zipfile
from collections import OrderedDict, defaultdict

import openpyxl

from padronizar_tag import (_ampliar_aba, _caminho_aba, _col_letra, _estilo, _gravar_celula,
                            _limpar, padronizar_tag)


def _norm(v):
    s = unicodedata.normalize("NFKD", str(v or "")).encode("ascii", "ignore").decode()
    s = s.upper().replace("º", "O").replace("°", "O")
    return re.sub(r"\s+", " ", s).strip()


# --------------------------------------------------------------------------
# Etapas
# --------------------------------------------------------------------------

# (chave, nome, grupos no Mapa de Juntas, cabeçalhos no Mapa de Spools)
ETAPAS = [
    ("VA", "Visual de Ajuste",
     {"VISUAL AJUSTE", "VISUAL DE AJUSTE", "VA"},
     ["DATA DE CORTE", "DATA CORTE", "DATA VA", "DATA DE VA", "VA", "VISUAL DE AJUSTE",
      "VISUAL AJUSTE", "DATA VISUAL DE AJUSTE", "DATA VISUAL AJUSTE"]),
    ("AJUSTE", "Ajuste",
     {"AJUSTE"},
     ["DATA DE AJUSTE", "DATA AJUSTE", "AJUSTE"]),
    ("SOLDAGEM", "Soldagem",
     {"SOLDAGEM", "SOLDA"},
     ["DATA DE SOLDAGEM", "DATA SOLDAGEM", "SOLDAGEM"]),
    ("EVS", "EVS",
     {"ENSAIO VISUAL", "EVS", "ENSAIO VISUAL DE SOLDA", "INSPECAO VISUAL DE SOLDA"},
     ["DATA DE EVS", "DATA EVS", "EVS"]),
    ("END", "END",
     {"LIQUIDO PENETRANTE / PM", "LIQUIDO PENETRANTE", "LP", "PM", "LP/PM", "PARTICULA MAGNETICA",
      "RX/US", "RX", "US", "ULTRASSOM", "RADIOGRAFIA", "END"},
     ["DATA DE END", "DATA END", "END"]),
    ("DIMENSIONAL", "Dimensional",
     {"DIMENSIONAL", "CONTROLE DIMENSIONAL", "INSPECAO DIMENSIONAL"},
     ["DATA DIMENSIONAL", "DATA DE DIMENSIONAL", "DATA DO DIMENSIONAL", "DIMENSIONAL"]),
]

# título das colunas criadas com --criar-colunas
TITULOS = {"VA": "Data de Corte", "AJUSTE": "Data de Ajuste", "SOLDAGEM": "Data de Soldagem",
           "EVS": "Data de EVS", "END": "Data de END", "DIMENSIONAL": "Data Dimensional"}

PIPE = {"P", "PIPE", "PIPE SHOP", "PIPESHOP", "OFICINA"}
CAMPO = {"C", "CAMPO", "CAMP", "FIELD"}

CAB_LINHA = ["LINHA", "NO DA LINHA", "N DA LINHA", "N. DA LINHA", "LINE", "LINE NUMBER", "TAG LINHA",
             "TAG DA LINHA", "TAG", "ISOMETRICO"]
CAB_SPOOL = ["SPOOL", "NO SPOOL", "NO DO SPOOL", "N SPOOL", "N DO SPOOL", "SPOOL NO"]


# --------------------------------------------------------------------------
# Datas
# --------------------------------------------------------------------------

def ler_data(v):
    """Retorna (date|None, preenchido). preenchido=True com date=None indica valor inválido."""
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None, False
    if isinstance(v, dt.datetime):
        return v.date(), True
    if isinstance(v, dt.date):
        return v, True
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 20000 < v < 80000:  # nº serial do Excel
            return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date(), True
        return None, True
    m = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})", str(v).strip())
    if not m:
        return None, True
    d, mes, a = (int(x) for x in m.groups())
    if a < 1000:  # 26 ou 026 (erro de digitação) -> 2026
        a = 2000 + a % 100
    try:
        return dt.date(a, mes, d), True
    except ValueError:
        return None, True


def _serial(d):
    return (d - dt.date(1899, 12, 30)).days


# --------------------------------------------------------------------------
# Leitura do Mapa de Juntas
# --------------------------------------------------------------------------

def _chave_linha(v):
    return _limpar(v).replace('"', "") if v not in (None, "") else ""


def _chave_tag(v):
    return padronizar_tag(v)[0].replace('"', "") if v not in (None, "") else ""


def _spool(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        v = int(v)
    s = str(v).strip().upper()
    return s.zfill(3) if s.isdigit() else s


def _junta(v):
    return _spool(v)


def localizar_juntas(ws, max_linhas=30):
    """Linha de cabeçalho (Linha/Spool/Junta/P/C) e colunas de data por etapa."""
    linhas = list(ws.iter_rows(min_row=1, max_row=max_linhas, values_only=True))
    for i, row in enumerate(linhas):
        hs = [_norm(c) for c in row]
        if "SPOOL" in hs and "P/C" in hs and "JUNTA" in hs:
            break
    else:
        raise SystemExit("Cabeçalho do Mapa de Juntas (Linha/Spool/Junta/P/C) não encontrado em %r" % ws.title)
    hdr = i + 1
    hs = [_norm(c) for c in linhas[i]]
    sub = [_norm(c) for c in linhas[i + 1]] if i + 1 < len(linhas) else []
    col = {}
    col["linha"] = hs.index("LINHA") + 1 if "LINHA" in hs else None
    col["iso"] = hs.index("ISOMETRICO") + 1 if "ISOMETRICO" in hs else None
    if not col["linha"] and not col["iso"]:
        raise SystemExit("Coluna 'Linha' não encontrada no Mapa de Juntas")
    col["spool"] = hs.index("SPOOL") + 1
    col["junta"] = hs.index("JUNTA") + 1
    col["pc"] = hs.index("P/C") + 1

    # grupos: cabeçalho na linha hdr, subcolunas (Status/Data...) na linha seguinte
    inicios = [(j, h) for j, h in enumerate(hs) if h]
    etapas = defaultdict(list)  # chave -> [(nome_grupo, col_status|None, [cols_data])]
    for n, (j, h) in enumerate(inicios):
        fim = inicios[n + 1][0] if n + 1 < len(inicios) else max(len(hs), len(sub))
        datas = [k + 1 for k in range(j, fim) if k < len(sub) and sub[k].startswith("DATA")]
        status = next((k + 1 for k in range(j, fim) if k < len(sub) and sub[k] == "STATUS"), None)
        if not datas:
            continue
        for chave, _nome, grupos, _ in ETAPAS:
            if h in grupos:
                etapas[chave].append((h, status, datas))
    tem_sub = any(s for s in sub)
    return hdr + (2 if tem_sub else 1), col, dict(etapas)


def ler_juntas(caminho, aba=None):
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    abas = [wb[aba]] if aba else [ws for ws in wb.worksheets if _tem_juntas(ws)]
    juntas, etapas_encontradas, avisos = [], set(), []
    for ws in abas:
        primeira, col, etapas = localizar_juntas(ws)
        etapas_encontradas |= set(etapas)
        for r, row in enumerate(ws.iter_rows(min_row=primeira, values_only=True), primeira):
            def v(c):
                return row[c - 1] if c and c - 1 < len(row) else None
            linha_bruta = v(col["linha"]) or v(col["iso"])
            spool = _spool(v(col["spool"]))
            if not linha_bruta or not spool:
                continue
            pc = _norm(v(col["pc"]))
            tipo = "PIPE" if pc in PIPE else "CAMPO" if pc in CAMPO else ""
            if not tipo and pc:
                avisos.append("%s!%d: P/C desconhecido %r (junta desconsiderada)" % (ws.title, r, v(col["pc"])))
            datas = {}
            for chave, grupos in etapas.items():
                datas[chave] = _data_etapa(chave, grupos, v)
            juntas.append({
                "aba": ws.title, "linha_xlsx": r, "linha": str(linha_bruta).strip(),
                "chave": _chave_linha(linha_bruta), "tag": _chave_tag(linha_bruta),
                "spool": spool, "junta": _junta(v(col["junta"])), "tipo": tipo, "datas": datas,
            })
    return juntas, etapas_encontradas, avisos


def _tem_juntas(ws):
    for row in ws.iter_rows(min_row=1, max_row=30, values_only=True):
        hs = [_norm(c) for c in row]
        if "SPOOL" in hs and "P/C" in hs and "JUNTA" in hs:
            return True
    return False


def _data_etapa(chave, grupos, v):
    """(date|None, motivo) da junta na etapa. motivo explica a ausência de data."""
    if chave != "END":
        datas, invalidas = [], []
        for _h, _st, cols in grupos:
            for c in cols:
                d, preenchido = ler_data(v(c))
                if d:
                    datas.append(d)
                elif preenchido:
                    invalidas.append(str(v(c)))
        if datas:
            return max(datas), ""
        return None, ("data inválida %r" % invalidas[0]) if invalidas else "sem data"
    # END: vários ensaios; "N" = não aplicável, "P" sem data = pendente
    datas, pendentes = [], []
    for h, st, cols in grupos:
        status = _norm(v(st)) if st else ""
        if status == "N":
            continue
        ds = [d for d in (ler_data(v(c))[0] for c in cols) if d]
        if ds:
            datas.extend(ds)
        elif status == "P":
            pendentes.append(h)
    if pendentes:
        return None, "END pendente (%s)" % ", ".join(pendentes)
    if datas:
        return max(datas), ""
    return None, "sem data"


# --------------------------------------------------------------------------
# Consolidação por spool
# --------------------------------------------------------------------------

def consolidar(juntas, etapas):
    """{(chave_linha, spool): {"juntas": [...], "etapas": {chave: (date|None, motivo)}}}"""
    spools = OrderedDict()
    for j in juntas:
        s = spools.setdefault((j["chave"], j["spool"]), {"linha": j["linha"], "tag": j["tag"],
                                                         "juntas": [], "campo": []})
        (s["juntas"] if j["tipo"] == "PIPE" else s["campo"]).append(j)
    for s in spools.values():
        s["etapas"] = {}
        for chave in etapas:
            if not s["juntas"]:
                s["etapas"][chave] = (None, "spool sem juntas Pipe Shop")
                continue
            faltando = [j for j in s["juntas"] if not j["datas"][chave][0]]
            if faltando:
                s["etapas"][chave] = (None, "%d de %d juntas Pipe sem data: %s" % (
                    len(faltando), len(s["juntas"]),
                    ", ".join("%s (%s)" % (j["junta"], j["datas"][chave][1]) for j in faltando[:10])
                    + (" ..." if len(faltando) > 10 else "")))
            else:
                s["etapas"][chave] = (max(j["datas"][chave][0] for j in s["juntas"]), "")
    return spools


# --------------------------------------------------------------------------
# Mapa de Spools
# --------------------------------------------------------------------------

def localizar_spools(ws, max_linhas=30):
    for row in ws.iter_rows(min_row=1, max_row=max_linhas):
        hs = {c.column: _norm(c.value) for c in row if c.value not in (None, "")}
        c_spool = next((c for c, h in hs.items() if h in CAB_SPOOL), None)
        c_linha = None
        for alvo in CAB_LINHA:  # ordem de preferência
            c_linha = next((c for c, h in hs.items() if h == alvo), None)
            if c_linha:
                break
        if c_spool and c_linha:
            destinos = {}
            for chave, _nome, _g, cabs in ETAPAS:
                for alvo in cabs:
                    c = next((c for c, h in hs.items() if h == alvo and c not in destinos.values()), None)
                    if c:
                        destinos[chave] = c
                        break
            return row[0].row, c_linha, c_spool, destinos, max(hs)
    raise SystemExit("Cabeçalho do Mapa de Spools (Linha + Spool) não encontrado em %r" % ws.title)


def _cellxfs(styles):
    m = re.search(r"<cellXfs[^>]*>(.*?)</cellXfs>", styles, re.S)
    return m, re.findall(r"<xf\b[^>]*?(?:/>|>.*?</xf>)", m.group(1), re.S)


def _numfmt_data(styles, codigo="dd/mm/yyyy"):
    """Id do formato de número dd/mm/yyyy (criado em <numFmts> se não existir)."""
    m = re.search(r'<numFmt numFmtId="(\d+)" formatCode="%s"' % re.escape(codigo), styles)
    if m:
        return int(m.group(1)), styles
    ids = [int(i) for i in re.findall(r'<numFmt numFmtId="(\d+)"', styles)]
    novo_id = max(ids + [163]) + 1
    tag = '<numFmt numFmtId="%d" formatCode="%s"/>' % (novo_id, codigo)
    novo = '<numFmts count="%d">' % (len(ids) + 1)
    if re.search(r"<numFmts\b[^>]*/>", styles):
        styles = re.sub(r"<numFmts\b[^>]*/>", novo + tag + "</numFmts>", styles, count=1)
    elif "<numFmts" in styles:
        styles = re.sub(r"<numFmts\b[^>]*>", novo, styles, count=1)
        styles = styles.replace("</numFmts>", tag + "</numFmts>", 1)
    else:
        styles = re.sub(r"(<styleSheet\b[^>]*>)", r'\1<numFmts count="1">%s</numFmts>' % tag, styles, count=1)
    return novo_id, styles


def _estilo_data(styles, base, cache):
    """Índice de um estilo igual a `base` com formato de data (dd/mm/aaaa)."""
    base = int(base or 0)
    if base in cache:
        return cache[base], styles
    m, xfs = _cellxfs(styles)
    xf = xfs[base] if base < len(xfs) else xfs[0]
    fmt = re.search(r'numFmtId="(\d+)"', xf)
    fmt_id = int(fmt.group(1)) if fmt else 0
    nfmt = re.search(r'<numFmt numFmtId="%d" formatCode="([^"]*)"' % fmt_id, styles)
    if 14 <= fmt_id <= 22 or (nfmt and re.search(r"[dy]", nfmt.group(1).lower()) and "h" not in nfmt.group(1).lower()):
        cache[base] = base
        return base, styles
    id_data, styles = _numfmt_data(styles)
    m, xfs = _cellxfs(styles)
    novo = re.sub(r'\snumFmtId="\d+"', "", xf)
    novo = re.sub(r'\sapplyNumberFormat="\d+"', "", novo)
    novo = re.sub(r"^<xf\b", '<xf numFmtId="%d" applyNumberFormat="1"' % id_data, novo)
    if novo in xfs:
        cache[base] = xfs.index(novo)
        return cache[base], styles
    corpo = m.group(1) + novo
    ini = styles[:m.start()] + re.sub(r'count="\d+"', 'count="%d"' % (len(xfs) + 1),
                                      styles[m.start():m.start(1)], count=1)
    styles = ini + corpo + styles[m.end(1):]
    cache[base] = len(xfs)
    return len(xfs), styles


def atualizar_spools(juntas_xlsx, spools_xlsx, saida, aba_juntas=None, aba_spools=None,
                     criar_colunas=False, limpar=True):
    juntas, etapas_juntas, avisos = ler_juntas(juntas_xlsx, aba_juntas)
    spools = consolidar(juntas, etapas_juntas)
    por_tag = defaultdict(set)
    for (chave, sp), s in spools.items():
        por_tag[(s["tag"], sp)].add(chave)

    wb = openpyxl.load_workbook(spools_xlsx)
    ws = wb[aba_spools] if aba_spools else wb.worksheets[0]
    hdr, c_linha, c_spool, destinos, ultima = localizar_spools(ws)

    novas = []
    for chave, nome, _g, cabs in ETAPAS:
        if chave not in etapas_juntas:
            avisos.append("Etapa %s: coluna de data não encontrada no Mapa de Juntas (ignorada)" % nome)
            continue
        if chave not in destinos:
            if criar_colunas:
                ultima += 1
                destinos[chave] = ultima
                novas.append((ultima, TITULOS[chave]))
            else:
                avisos.append("Etapa %s: coluna %r não encontrada no Mapa de Spools (use --criar-colunas)"
                              % (nome, cabs[0]))
    destinos = {k: c for k, c in destinos.items() if k in etapas_juntas}

    gravar, resultado, sem_juntas = {}, [], []
    for r in range(hdr + 1, ws.max_row + 1):
        linha, sp = ws.cell(r, c_linha).value, _spool(ws.cell(r, c_spool).value)
        if not linha or not sp:
            continue
        s = spools.get((_chave_linha(linha), sp))
        if s is None:
            cand = por_tag.get((_chave_tag(linha), sp), set())
            if len(cand) == 1:
                s = spools[(next(iter(cand)), sp)]
        if s is None:
            sem_juntas.append((r, linha, sp))
            continue
        for chave in (k for k, *_ in ETAPAS if k in etapas_juntas):
            data, motivo = s["etapas"][chave]
            c = destinos.get(chave)
            if c:
                atual, _ = ler_data(ws.cell(r, c).value)
                ref = "%s%d" % (_col_letra(c), r)
                if data and data != atual:
                    gravar[ref] = data
                elif not data and ws.cell(r, c).value not in (None, "") and limpar:
                    gravar[ref] = None
            resultado.append((r, linha, sp, chave, data, motivo, len(s["juntas"]), len(s["campo"])))

    with zipfile.ZipFile(spools_xlsx) as zin:
        caminho = _caminho_aba(zin, ws.title)
        xml = zin.read(caminho).decode("utf-8")
        styles = zin.read("xl/styles.xml").decode("utf-8")
        cache = {}
        for c, titulo in novas:
            ant = _col_letra(c - 1)
            xml = _gravar_celula(xml, "%s%d" % (_col_letra(c), hdr), titulo, _estilo(xml, "%s%d" % (ant, hdr)))
            xml = _ampliar_aba(xml, c)
        for ref, data in gravar.items():
            if data is None:
                xml = _gravar_celula(xml, ref, None)
                continue
            base = _estilo(xml, ref)
            if base is None:  # célula nova: herda o estilo da coluna anterior na mesma linha
                lin = re.match(r"[A-Z]+(\d+)", ref).group(1)
                col = re.match(r"[A-Z]+", ref).group(0)
                n = sum((ord(ch) - 64) * 26 ** i for i, ch in enumerate(reversed(col)))
                base = _estilo(xml, "%s%s" % (_col_letra(n - 1), lin)) if n > 1 else None
            estilo, styles = _estilo_data(styles, base, cache)
            xml = _gravar_celula(xml, ref, None, estilo, numero=_serial(data), forcar_estilo=True)
        alterados = {caminho: xml, "xl/styles.xml": styles}
        with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                dados = alterados[item.filename].encode("utf-8") if item.filename in alterados \
                    else zin.read(item.filename)
                zout.writestr(item, dados)

    return {"aba": ws.title, "destinos": destinos, "gravados": gravar, "resultado": resultado,
            "sem_juntas": sem_juntas, "avisos": avisos, "spools": spools, "juntas": juntas}


def salvar_relatorio(caminho, res):
    nomes = {k: n for k, n, _g, _c in ETAPAS}
    with open(caminho, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Linha planilha", "Linha", "Spool", "Etapa", "Data consolidada", "Juntas Pipe",
                    "Juntas Campo (desconsideradas)", "Pendência"])
        for r, linha, sp, chave, data, motivo, n_pipe, n_campo in res["resultado"]:
            w.writerow([r, linha, sp, nomes[chave], data.strftime("%d/%m/%Y") if data else "",
                        n_pipe, n_campo, motivo])
        for r, linha, sp in res["sem_juntas"]:
            w.writerow([r, linha, sp, "", "", 0, 0, "spool não encontrado no Mapa de Juntas"])


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("mapa_juntas")
    ap.add_argument("mapa_spools")
    ap.add_argument("saida")
    ap.add_argument("--aba-juntas", help="aba do Mapa de Juntas (padrão: todas as abas com juntas)")
    ap.add_argument("--aba-spools", help="aba do Mapa de Spools (padrão: primeira)")
    ap.add_argument("--criar-colunas", action="store_true",
                    help="cria no Mapa de Spools as colunas de data que não existirem")
    ap.add_argument("--nao-limpar", action="store_true",
                    help="não apaga datas já existentes de etapas que ficaram incompletas")
    ap.add_argument("--relatorio", help="CSV com a situação de cada spool/etapa")
    a = ap.parse_args(argv)

    res = atualizar_spools(a.mapa_juntas, a.mapa_spools, a.saida, a.aba_juntas, a.aba_spools,
                           a.criar_colunas, not a.nao_limpar)
    nomes = {k: n for k, n, _g, _c in ETAPAS}
    n_pipe = sum(1 for j in res["juntas"] if j["tipo"] == "PIPE")
    print("Mapa de Juntas: %d juntas (%d Pipe Shop), %d spools" % (len(res["juntas"]), n_pipe, len(res["spools"])))
    print("Aba %s: colunas atualizadas: %s" % (res["aba"], ", ".join(
        "%s (%s)" % (nomes[k], _col_letra(c)) for k, c in res["destinos"].items()) or "nenhuma"))
    preench = defaultdict(int)
    for _r, _l, _s, chave, data, *_ in res["resultado"]:
        if data:
            preench[chave] += 1
    for k, n, *_ in ETAPAS:
        if k in res["destinos"]:
            print("  %-16s %d spools concluídos" % (n, preench[k]))
    print("%d células alteradas" % len(res["gravados"]))
    for r, linha, sp in res["sem_juntas"]:
        print("  linha %d: spool %s / %s não encontrado no Mapa de Juntas" % (r, sp, linha))
    for av in res["avisos"]:
        print("  AVISO:", av)
    if a.relatorio:
        salvar_relatorio(a.relatorio, res)
        print("Relatório: %s" % a.relatorio)
    return 0


if __name__ == "__main__":
    sys.exit(main())
