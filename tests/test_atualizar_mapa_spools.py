"""Testes da consolidação Mapa de Juntas -> Mapa de Spools.

    python -m pytest tests/      (ou: python tests/test_atualizar_mapa_spools.py)
"""
import datetime as dt
import os
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))
from atualizar_mapa_spools import atualizar_spools, ler_data  # noqa: E402

D = dt.datetime
LINHA = '1"-A3B-N2-9114-Z01'


def _mapa_juntas(caminho, juntas):
    """Layout do ControlTub: cabeçalho de grupos + linha de subcolunas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SGJ"
    ws.append(["SITUAÇÃO GERAL DE JUNTAS"])
    ws.append(["Linha", "Spool", "Junta", "P/C", "Visual Ajuste", None, "Soldagem", None,
               "Ensaio Visual", None, "Líquido Penetrante / PM", None, None, "RX/US", None, None])
    ws.append([None, None, None, None, "Status", "Data", "Status", "Data", "Status", "Data",
               "Status", "Data / LP", "Data / PM", "Status", "Data RX", "Data US"])
    for j in juntas:
        ws.append(j)
    wb.save(caminho)


def _mapa_spools(caminho, linhas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Linha", "Spool", "Data de Corte", "Data de Soldagem", "Data de EVS", "Data de END"])
    for l in linhas:
        ws.append(l)
    wb.save(caminho)


def _rodar(juntas, spools, **kw):
    d = tempfile.mkdtemp()
    mj, ms, out = (os.path.join(d, n) for n in ("j.xlsx", "s.xlsx", "o.xlsx"))
    _mapa_juntas(mj, juntas)
    _mapa_spools(ms, spools)
    res = atualizar_spools(mj, ms, out, **kw)
    ws = openpyxl.load_workbook(out).active
    return res, [[c.value for c in row] for row in ws.iter_rows(min_row=2)], ws


def _j(junta, pc, va=None, sold=None, evs=None, lp=None, rx_st=None, rx=None, lp_st=None):
    return [LINHA, "001", junta, pc, "A" if va else "P", va, "A" if sold else "P", sold,
            "A" if evs else "P", evs, lp_st or ("A" if lp else "P"), lp, None, rx_st, rx, None]


def test_exemplo_do_procedimento_campo_desconsiderado():
    # juntas 001-004 Pipe Shop com VA; 005 de Campo sem VA -> spool concluído no VA
    juntas = [
        _j("001", "P", va=D(2026, 8, 10)),
        _j("002", "PIPE", va=D(2026, 8, 12)),
        _j("003", "P", va="15/08/26"),
        _j("004", "p", va=D(2026, 8, 11)),
        _j("005", "C"),
    ]
    _, linhas, ws = _rodar(juntas, [[LINHA, "001"]])
    assert linhas[0][2] == D(2026, 8, 15)          # maior data entre 001-004
    assert linhas[0][3] is None                    # soldagem sem datas
    assert ws.cell(2, 3).number_format == "dd/mm/yyyy"


def test_uma_junta_pipe_sem_data_deixa_em_branco_e_limpa_valor_antigo():
    juntas = [
        _j("001", "P", va=D(2026, 8, 10), sold=D(2026, 8, 20)),
        _j("002", "P", va=D(2026, 8, 12)),
    ]
    _, linhas, _ = _rodar(juntas, [[LINHA, 1, None, D(2020, 1, 1)]])
    assert linhas[0][2] == D(2026, 8, 12)
    assert linhas[0][3] is None                    # 002 sem soldagem -> apaga data antiga


def test_nao_limpar_preserva_valor_existente():
    juntas = [_j("001", "P"), _j("002", "P", sold=D(2026, 8, 20))]
    _, linhas, _ = _rodar(juntas, [[LINHA, "001", None, D(2020, 1, 1)]], limpar=False)
    assert linhas[0][3] == D(2020, 1, 1)


def test_end_ignora_nao_aplicavel_e_bloqueia_pendente():
    juntas = [
        _j("001", "P", lp=D(2026, 8, 18), rx_st="N"),
        _j("002", "P", lp=D(2026, 8, 17), rx_st="A", rx=D(2026, 8, 21)),
    ]
    _, linhas, _ = _rodar(juntas, [[LINHA, "001"]])
    assert linhas[0][5] == D(2026, 8, 21)

    juntas[0] = _j("001", "P", lp=D(2026, 8, 18), rx_st="P")  # RX pendente
    _, linhas, _ = _rodar(juntas, [[LINHA, "001"]])
    assert linhas[0][5] is None


def test_linha_sem_zona_usa_tag_padronizada():
    juntas = [_j("001", "P", va=D(2026, 8, 10))]
    _, linhas, _ = _rodar(juntas, [['1"-A3B-N2-9114', "001"]])
    assert linhas[0][2] == D(2026, 8, 10)


def test_spool_inexistente_e_reportado():
    res, linhas, _ = _rodar([_j("001", "P", va=D(2026, 8, 10))], [[LINHA, "002"]])
    assert [x[:3] for x in res["sem_juntas"]] == [(2, LINHA, "002")]
    assert linhas[0][2] is None


def test_layout_controltub_documento_e_colunas_de_montagem():
    """Cabeçalhos reais do SGS; mesmo spool em dois documentos; Montagem intocada."""
    d = tempfile.mkdtemp()
    mj, ms, out = (os.path.join(d, n) for n in ("j.xlsx", "s.xlsx", "o.xlsx"))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Linha", "Documento", "Spool", "Junta", "P/C", "Visual Ajuste", None])
    ws.append([None, None, None, None, None, "Status", "Data"])
    ws.append([LINHA, "ISO-1", "003", "001", "P", "A", D(2026, 8, 10)])
    ws.append([LINHA, "ISO-1", "003", "001", "P", "A", D(2026, 8, 10)])  # repetida
    ws.append([LINHA, "ISO-2", "003", "002", "P", "P", None])
    wb.save(mj)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Linha", "Documento", "Spool", "Total_Juntas", "Data\nCorte", "Data\nVA Fab",
               "Data\nVA Mon"])
    ws.append([LINHA, "ISO-1", "003", 1, None, None, D(2020, 1, 1)])
    ws.append([LINHA, "ISO-2", "003", 3, D(2020, 1, 1), None, None])
    wb.save(ms)
    res = atualizar_spools(mj, ms, out)
    linhas = [[c.value for c in r] for r in openpyxl.load_workbook(out).active.iter_rows(min_row=2)]
    assert linhas[0][4:] == [D(2026, 8, 10), D(2026, 8, 10), D(2020, 1, 1)]
    assert linhas[1][4:] == [None, None, None]
    assert [x[:3] + x[4:] for x in res["divergencias"]] == [(3, LINHA, "003", 3, 1)]


def test_dimensional_pela_rastreabilidade_e_linha_de_numeracao():
    d = tempfile.mkdtemp()
    mj, ms, out = (os.path.join(d, n) for n in ("j.xlsx", "s.xlsx", "o.xlsx"))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Linha", "Spool", "Junta", "P/C", "Rastreabilidade - dimensional", None, None, None])
    ws.append([None, None, None, None, "Status", "Relatório", "Executante", "Data"])
    ws.append([1, 2, 3, 4, 5, 6, 7, 8])  # numeração das colunas (não é junta)
    ws.append([LINHA, "001", "001", "P", "A", "053/26", "X", D(2026, 9, 2)])
    ws.append([LINHA, "001", "002", "P", "A", "053/26", "X", D(2026, 9, 3)])
    ws.append([LINHA, "001", "003", "C", "P", None, None, None])
    wb.save(mj)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Linha", "Spool", "Data\nDF Fab", "Relatório DF"])
    ws.append([LINHA, "001", None, "053/26"])
    wb.save(ms)
    res = atualizar_spools(mj, ms, out)
    ws = openpyxl.load_workbook(out).active
    assert ws["C2"].value == D(2026, 9, 3) and ws["D2"].value == "053/26"
    assert len(res["juntas"]) == 3


def test_espessura_do_spool_e_a_maior_da_linha():
    d = tempfile.mkdtemp()
    mj, ms, out = (os.path.join(d, n) for n in ("j.xlsx", "s.xlsx", "o.xlsx"))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Linha", "Spool", "Junta", "Espessura", "P/C", "Visual Ajuste", None])
    ws.append([None, None, None, None, None, "Status", "Data"])
    ws.append([LINHA, "001", "001", " 3,05", "P", "P", None])
    ws.append([LINHA, "002", "002", 2.77, "P", "P", None])
    ws.append([LINHA, "002", "003", "3,76", "C", "P", None])  # Campo também conta
    ws.append(['2"-OUTRA-LINHA-0001', "001", "001", "5,49", "P", "P", None])
    wb.save(mj)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Linha", "Spool", "Espessura", "Data\nCorte"])
    ws.append([LINHA, "001", None, None])
    ws.append([LINHA, "002", None, None])
    ws.append(['2"-OUTRA-LINHA-0001', "001", None, None])
    wb.save(ms)
    atualizar_spools(mj, ms, out)
    ws = openpyxl.load_workbook(out).active
    assert [ws.cell(r, 3).value for r in (2, 3, 4)] == [3.76, 3.76, 5.49]


def test_ler_data():
    assert ler_data("10/09/026") == (dt.date(2026, 9, 10), True)
    assert ler_data("18/08/2026") == (dt.date(2026, 8, 18), True)
    assert ler_data("A") == (None, True)
    assert ler_data("") == (None, False)


if __name__ == "__main__":
    for nome, f in list(globals().items()):
        if nome.startswith("test_"):
            f()
            print("ok", nome)
